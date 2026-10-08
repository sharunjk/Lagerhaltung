"""Excel-Export und -Import (openpyxl)."""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KOPF_FILL = PatternFill("solid", fgColor="1F3A5F")
KOPF_FONT = Font(color="FFFFFF", bold=True)


def tabelle_xlsx(titel: str, spalten: list[str], zeilen: list[list], breiten: list[int] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = titel[:31]
    ws.append(spalten)
    for c in ws[1]:
        c.fill, c.font = KOPF_FILL, KOPF_FONT
        c.alignment = Alignment(vertical="center")
    for z in zeilen:
        ws.append([_xl(v) for v in z])
    for i, sp in enumerate(spalten, 1):
        w = (breiten[i - 1] if breiten and i - 1 < len(breiten) else None) or min(50, max(10, len(sp) + 2,
              *(len(str(z[i - 1])) + 2 for z in zeilen[:300] if i - 1 < len(z) and z[i - 1] is not None)))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xl(v):
    if isinstance(v, (int, float, date, datetime)) or v is None:
        return v
    try:
        from decimal import Decimal
        if isinstance(v, Decimal):
            return float(v)
    except Exception:
        pass
    return str(v)


IMPORT_SPALTEN = {
    "Artikelnummer": "nummer", "Bezeichnung": "bezeichnung", "Typenbezeichnung": "typ", "Gruppe": "kategorie",
    "Meldebestand": "meldebestand", "Mindestbestand": "mindestbestand", "Bestellmenge": "bestellmenge", "Einheit": "einheit",
    "Lieferant": "lieferant", "Lieferanten-Artikelnr": "lieferant_artnr", "Hersteller": "hersteller", "Hersteller-Nr": "hersteller_nr",
    "Preis": "preis", "Kritisch": "kritisch", "Verwendung": "verwendung", "Notiz": "notiz", "Lagerplatz": "lagerplatz",
    "Anfangsbestand": "anfangsbestand",
}


def import_vorlage() -> bytes:
    beispiel = ["", "Kugelhahn DN15 Edelstahl", "SS-45S8", "Kugelhahn", "2", "1", "4", "Stk", "Swagelok", "SS-45S8", "Swagelok",
                "SS-45S8", "125.50", "ja", "VV 1040 80", "", "C1-R2-3", "4"]
    return tabelle_xlsx("Artikelimport", list(IMPORT_SPALTEN), [beispiel])


def import_lesen(daten: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(daten), data_only=True, read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    if not rows:
        return []
    alias = {k.lower(): v for k, v in IMPORT_SPALTEN.items()}
    alias.update({"lagerort": "lagerplatz", "freifeld1": "lieferant", "freifeld2": "kategorie", "freifeld3": "typ", "kategorie": "kategorie"})
    kopf = [alias.get(str(h or "").strip().lower()) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(v not in (None, "") for v in r):
            continue
        d = {}
        for k, v in zip(kopf, r):
            if k:
                d[k] = "" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() and k != "preis" else str(v)).strip()
        out.append(d)
    return out
